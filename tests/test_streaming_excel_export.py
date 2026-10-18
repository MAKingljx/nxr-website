"""Exercise export artifacts and HTTP contracts using only disposable SQLite data."""
import io
import json
import os
import sqlite3
import tempfile
import unittest
from datetime import datetime
from pathlib import Path
from unittest import mock

_IMPORT_TEMP = tempfile.TemporaryDirectory()
_IMPORT_ROOT = Path(_IMPORT_TEMP.name)
for _key, _path in {
    'NXR_DATA_DIR': 'data', 'NXR_DB_PATH': 'data/cards.db',
    'NXR_TEMP_DB_PATH': 'data/temp_cards.db', 'NXR_SITE_STATIC_DIR': 'static',
    'NXR_ADMIN_UPLOAD_FOLDER': 'uploads',
}.items():
    os.environ[_key] = str(_IMPORT_ROOT / _path)
os.environ['NXR_DB_BACKEND'] = 'sqlite'
os.environ['NXR_SKIP_DB_INIT'] = '1'

import pandas as pd
from openpyxl import load_workbook
from nxr_admin import admin_core, routes_auth, routes_exports


class FixedDateTime(datetime):
    @classmethod
    def now(cls, tz=None):
        return cls(2026, 10, 18, 12, 34, 56)


def apply_score_number_format(worksheet, column_names=routes_exports.SCORE_EXPORT_COLUMNS):
    from openpyxl.utils import get_column_letter

    header_map = {
        cell.value: cell.column
        for cell in worksheet[1]
        if cell.value
    }
    for column_name in column_names:
        column_index = header_map.get(column_name)
        if not column_index:
            continue
        column_letter = get_column_letter(column_index)
        for cell in worksheet[column_letter][1:]:
            if isinstance(cell.value, (int, float)):
                cell.number_format = '0.0'


def legacy_workbook(conn, query, params, target, selection, cert_ids):
    """Pre-change export algorithm retained as an independent artifact oracle."""
    frame = pd.read_sql_query(query, conn, params=params)
    if frame.empty:
        return 0
    frame['landing_page_url'] = frame['cert_id'].apply(lambda value: f'nxrgrading.com/card/{str(value).strip()}')
    frame['final_grade_text'] = frame['final_grade_text'].apply(admin_core.normalize_final_grade_text).replace('', pd.NA).fillna(frame['final_grade_text'])
    routes_exports.normalize_score_columns_for_export(frame, pd)
    frame = frame[[column for column in frame.columns if column != 'landing_page_url'] + ['landing_page_url']]
    with pd.ExcelWriter(target, engine='openpyxl') as writer:
        frame.to_excel(writer, sheet_name='Approved Cards', index=False)
        apply_score_number_format(writer.sheets['Approved Cards'])
        pd.DataFrame({
            '导出时间': [FixedDateTime.now().strftime('%Y-%m-%d %H:%M:%S')],
            '总记录数': [len(frame)],
            '筛选条件': [routes_exports.format_export_filter_label(export_selection=selection, cert_ids=cert_ids)],
            '数据范围': [admin_core.format_export_date_range(frame)],
            '包含字段数': [len(frame.columns)], '文件名称': [target.name],
        }).to_excel(writer, sheet_name='导出汇总', index=False)
        stats = frame['final_grade_text'].value_counts().reset_index()
        stats.columns = ['评分等级', '数量']
        stats['占比'] = (stats['数量'] / len(frame) * 100).round(1).astype(str) + '%'
        stats.to_excel(writer, sheet_name='评分统计', index=False)
        apply_score_number_format(writer.sheets['评分统计'], column_names=('评分等级',))
    return len(frame)


def stream_workbook(conn, query, params, target, selection, cert_ids):
    with tempfile.TemporaryFile() as spool:
        routes_exports.spool_export_query(conn, query, params, spool)
        return routes_exports.write_streaming_export(spool, target, selection, cert_ids)


def workbook_contract(source):
    workbook = load_workbook(io.BytesIO(Path(source).read_bytes()))
    try:
        return [(sheet.title, [[(cell.value, cell.data_type, cell.number_format,
                                  cell.font.bold, cell.alignment.horizontal,
                                  cell.alignment.vertical, cell.border.bottom.style)
                                 for cell in row] for row in sheet.iter_rows()])
                for sheet in workbook.worksheets]
    finally:
        workbook.close()


class StreamingExcelExportTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name)
        self.database = self.root / 'cards.sqlite'
        self.conn = sqlite3.connect(self.database)
        self.conn.row_factory = sqlite3.Row
        self.addCleanup(self.conn.close)
        self.conn.execute('''CREATE TABLE temp_cards (
            id INTEGER PRIMARY KEY, cert_id TEXT, card_name TEXT,
            product_type TEXT, vintage_classification TEXT, card_category TEXT,
            brand TEXT, final_grade_text TEXT, centering REAL, final_grade REAL,
            ai_grade TEXT, ai_centering REAL, ai_edges REAL, ai_corners REAL, ai_surface REAL,
            entry_notes TEXT, approved_at TEXT, entry_date TEXT, created_at TEXT,
            updated_at TEXT, approval_sequence INTEGER, status TEXT, optional_number INTEGER
        )''')
        self.records = []
        kinds = ['graded_card', 'merch_product', 'vintage_product', 'graded_card',
                 'graded_card', 'label_product', None, 'graded_card', 'graded_card']
        grades = ['9.00', '10', 'Pristine 10', 'pristine-10', '8.50', '9', 'custom', '', '10']
        for i, (kind, grade) in enumerate(zip(kinds, grades)):
            self.records.append((i + 1, str(1234567800 + i), f'卡片 {i}', kind,
                                 'Nova' if kind == 'vintage_product' else '', 'trading_card',
                                 'Pokemon', grade, 9.25, 9.95, '9.35', 9.25, 8.75, 9.85, None,
                                 '=SUM(1,2)' if i == 0 else '',
                                 '2026-10-18T01:02:03' if i > 3 else None,
                                 f'2026-10-{10+i:02d}', '2026-09-01', '',
                                 i, 'approved', 4 if i == 0 else None))
        self.conn.executemany('INSERT INTO temp_cards VALUES (' + ','.join(['?'] * 23) + ')', self.records)
        self.conn.commit()
        self.client = admin_core.app.test_client()
        admin_core.app.config.update(TESTING=True, SECRET_KEY='test-only-disposable-session')
        self.patches = [
            mock.patch.object(routes_exports, 'get_temp_db_connection', self.connect),
            mock.patch.object(routes_exports, 'ADMIN_DIR', self.root),
            mock.patch.object(admin_core, 'ADMIN_DIR', self.root),
            mock.patch.object(routes_exports, 'EXPORT_CHUNK_SIZE', 2),
            mock.patch.object(routes_exports, 'datetime', FixedDateTime),
        ]
        for patch in self.patches:
            patch.start()
            self.addCleanup(patch.stop)
        with self.client.session_transaction() as session:
            session.update(admin_logged_in=True, username='fixture', role='superadmin')

    def connect(self):
        conn = sqlite3.connect(self.database)
        conn.row_factory = sqlite3.Row
        return conn

    def query(self, filter_value='all', cert_ids=None):
        selection = routes_exports.parse_export_selection(filter_value)
        query, params = routes_exports.build_export_filter_sql(cert_ids=cert_ids, **routes_exports.export_selection_query_filters(selection))
        return selection, query + ' ORDER BY ' + admin_core.build_approved_order_clause(), params

    def test_workbooks_match_legacy_for_every_filter_and_chunk_boundary(self):
        cases = [('all', None), ('grade:9', None), ('8.5', None), ('grade:Pristine 10', None),
                 ('merch_product', None), ('vintage_product:Nova', None),
                 ('all', ['1234567800', '1234567803', '1234567804'])]
        for filter_value, cert_ids in cases:
            with self.subTest(filter=filter_value, cert_ids=cert_ids):
                old_dir, new_dir = self.root / 'old', self.root / 'new'
                old_dir.mkdir(exist_ok=True)
                new_dir.mkdir(exist_ok=True)
                old_path, new_path = old_dir / 'same.xlsx', new_dir / 'same.xlsx'
                selection, query, params = self.query(filter_value, cert_ids)
                expected = legacy_workbook(self.conn, query, params, old_path, selection, cert_ids or [])
                actual = stream_workbook(self.conn, query, params, new_path, selection, cert_ids or [])
                self.assertEqual(expected, actual)
                self.assertEqual(workbook_contract(old_path), workbook_contract(new_path))

    def test_real_generate_list_download_and_preview_readback(self):
        before = list(self.conn.execute('SELECT * FROM temp_cards ORDER BY id'))
        preview = self.client.post('/admin/export/preview', data={'export_filter': 'all'})
        self.assertEqual(preview.status_code, 200)
        self.assertEqual(preview.json['total_count'], 9)
        response = self.client.post('/admin/export/generate-excel', data={'export_filter': 'all'})
        self.assertEqual(response.status_code, 302)
        listing = self.client.get('/admin/export/list').json
        self.assertEqual(listing['total'], 1)
        filename = listing['exports'][0]['name']
        downloaded = self.client.get('/admin/export/download/' + filename)
        self.assertEqual(downloaded.status_code, 200)
        self.assertIn('attachment;', downloaded.headers['Content-Disposition'])
        workbook = load_workbook(io.BytesIO(downloaded.data))
        self.assertEqual(workbook.sheetnames, ['Approved Cards', '导出汇总', '评分统计'])
        values = list(workbook['Approved Cards'].values)
        self.assertEqual(len(values), 10)
        self.assertEqual(values[0][-1], 'landing_page_url')
        rows = {row[1]: dict(zip(values[0], row)) for row in values[1:]}
        self.assertEqual(rows['1234567800']['final_grade_text'], 9)
        self.assertEqual(rows['1234567800']['centering'], 9.2)
        self.assertEqual(rows['1234567800']['card_name'], '卡片 0')
        self.assertIsNone(rows['1234567801']['final_grade_text'])
        self.assertIsNone(rows['1234567802']['ai_grade'])
        self.assertEqual(rows['1234567803']['final_grade_text'], 'Pristine 10')
        self.assertEqual(workbook['导出汇总']['B2'].value, 9)
        workbook.close()
        downloaded.close()
        history = json.loads((self.root / 'exports/export_history.json').read_text())
        self.assertEqual(history[-1]['record_count'], 9)
        self.assertEqual(history[-1]['file_size'], len(downloaded.data))
        self.assertEqual(before, list(self.conn.execute('SELECT * FROM temp_cards ORDER BY id')))

    def test_invalid_missing_empty_and_anonymous_requests_create_no_files(self):
        for form in ({'cert_ids': 'bad'}, {'cert_ids': '9999999999'}, {'export_filter': 'grade:9.5'}, {'export_filter': 'bad'}):
            with self.subTest(form=form):
                self.assertEqual(self.client.post('/admin/export/generate-excel', data=form).status_code, 302)
                self.assertEqual(list((self.root / 'exports').glob('*.xlsx')), [])
                self.assertFalse((self.root / 'exports/export_history.json').exists())
        anonymous = admin_core.app.test_client()
        for path in ['/admin/export/generate-excel', '/admin/export/download/fake.xlsx']:
            response = anonymous.post(path) if path.endswith('generate-excel') else anonymous.get(path)
            self.assertEqual(response.status_code, 302)
            self.assertIn('login', response.location)

    def test_delete_journal_writer_is_unblocked_during_excel_generation(self):
        self.assertEqual(self.conn.execute('PRAGMA journal_mode').fetchone()[0], 'delete')
        original_writer = routes_exports.write_streaming_export
        observed = []
        def write_with_concurrent_insert(*args, **kwargs):
            # The HTTP request is between SQL snapshot and long workbook IO.
            # A separate connection must be able to commit a real new row now.
            with sqlite3.connect(self.database, timeout=0.05) as other:
                other.execute("INSERT INTO temp_cards (id, cert_id, status) VALUES (99, '9999999999', 'approved')")
                other.commit()
            observed.append(True)
            return original_writer(*args, **kwargs)
        with mock.patch.object(routes_exports, 'write_streaming_export', side_effect=write_with_concurrent_insert):
            response = self.client.post('/admin/export/generate-excel', data={'export_filter': 'all'})
        self.assertEqual(response.status_code, 302)
        self.assertEqual(observed, [True])
        self.assertEqual(self.conn.execute('SELECT COUNT(*) FROM temp_cards').fetchone()[0], 10)
        history = json.loads((self.root / 'exports/export_history.json').read_text())
        self.assertEqual(history[-1]['record_count'], 9)
        self.assertEqual(self.conn.execute('PRAGMA journal_mode').fetchone()[0], 'delete')

    def test_statistics_ties_date_fallback_and_save_failure(self):
        self.conn.execute('DELETE FROM temp_cards')
        for i in range(25):
            self.conn.execute(
                "INSERT INTO temp_cards (id, cert_id, product_type, final_grade_text, entry_date, status) VALUES (?, ?, 'graded_card', ?, ?, 'approved')",
                (i + 1, str(1234567000 + i), f'Custom {i % 19}', ' 2026-08-01 ' if i else '2026-09-01'),
            )
        self.conn.commit()
        selection, query, params = self.query()
        old_path, new_path = self.root / 'old/same.xlsx', self.root / 'new/same.xlsx'
        old_path.parent.mkdir()
        new_path.parent.mkdir()
        legacy_workbook(self.conn, query, params, old_path, selection, [])
        stream_workbook(self.conn, query, params, new_path, selection, [])
        self.assertEqual(workbook_contract(old_path), workbook_contract(new_path))
        previous = new_path.read_bytes()
        with mock.patch.object(routes_exports.os, 'replace', side_effect=OSError('injected disk error')):
            with self.assertRaisesRegex(OSError, 'injected disk error'):
                stream_workbook(self.conn, query, params, new_path, selection, [])
        self.assertEqual(previous, new_path.read_bytes())
        self.assertEqual(list(new_path.parent.glob('.export-*.tmp')), [])

    def test_sql_reads_are_chunked_and_failure_leaves_no_export(self):
        selection, query, params = self.query()
        original_read = pd.read_sql_query
        with mock.patch.object(pd, 'read_sql_query', wraps=original_read) as read:
            stream_workbook(self.conn, query, params, self.root / 'bounded.xlsx', selection, [])
            self.assertEqual(read.call_args.kwargs['chunksize'], 2)
        normalizer = routes_exports.normalize_score_columns_for_export
        calls = 0
        def fail_after_one_chunk(frame, pd_module):
            nonlocal calls
            calls += 1
            if calls == 2:
                raise RuntimeError('injected write error')
            return normalizer(frame, pd_module)
        from openpyxl.worksheet._writer import ALL_TEMP_FILES
        before_spools = list(ALL_TEMP_FILES)
        with mock.patch.object(routes_exports, 'normalize_score_columns_for_export', side_effect=fail_after_one_chunk):
            self.client.post('/admin/export/generate-excel', data={'export_filter': 'all'})
        self.assertEqual(before_spools, ALL_TEMP_FILES)
        self.assertEqual(list((self.root / 'exports').glob('*')), [])


if __name__ == '__main__':
    unittest.main()
